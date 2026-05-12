#!/usr/bin/env python
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation


EMU_PER_INCH = 914400


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect a PPTX for basic deck QA.")
    parser.add_argument("pptx", type=Path)
    args = parser.parse_args()

    prs = Presentation(args.pptx)
    width = prs.slide_width / EMU_PER_INCH
    height = prs.slide_height / EMU_PER_INCH
    print(f"path: {args.pptx.resolve()}")
    print(f"slides: {len(prs.slides)}")
    print(f"size_in: {width:.2f} x {height:.2f}")
    print(f"aspect: {width / height:.3f}")

    for i, slide in enumerate(prs.slides, 1):
        text_shapes = 0
        picture_shapes = 0
        text_chars = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                txt = shape.text or ""
                if txt.strip():
                    text_shapes += 1
                    text_chars += len(txt)
            if shape.shape_type == 13:
                picture_shapes += 1
        print(
            f"slide {i:02d}: text_shapes={text_shapes}, "
            f"text_chars={text_chars}, pictures={picture_shapes}"
        )


if __name__ == "__main__":
    main()
